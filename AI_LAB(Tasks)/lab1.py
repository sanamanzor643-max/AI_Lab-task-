from docx import Document

doc = Document()

doc.add_heading("Artificial Intelligence Timeline", 0)

milestones = [
    ("1950 - Turing Test",
     "Alan Turing proposed a test to discuss machine intelligence.",
     "https://plato.stanford.edu/archives/spr2024/entries/artificial-intelligence/"),

    ("1956 - Dartmouth Workshop",
     "This workshop helped establish AI as a research field.",
     "https://ai100.stanford.edu/2016-report/appendix-i-short-history-ai"),

    ("1958 - Perceptron",
     "Frank Rosenblatt developed an early trainable neural network.",
     "https://www.ibm.com/think/topics/history-of-artificial-intelligence"),

    ("1974-1980 - AI Winter",
     "AI research faced reduced funding and interest because expectations were not met.",
     "https://ai100.stanford.edu/2016-report/appendix-i-short-history-ai"),

    ("1980s - Expert Systems",
     "Expert systems used rules to solve problems in specialized fields.",
     "https://ai100.stanford.edu/2016-report/appendix-i-short-history-ai"),

    ("2012 - AlexNet and Deep Learning",
     "GPU computing and large datasets helped deep learning achieve major image-recognition results.",
     "https://spectrum.ieee.org/alexnet-source-code"),

    ("2017 - Transformer",
     "The Transformer architecture introduced attention-based methods important to modern AI.",
     "https://arxiv.org/abs/1706.03762"),

    ("2022 - ChatGPT",
     "ChatGPT made conversational generative AI widely accessible.",
     "https://openai.com/index/chatgpt/")
]

for year, event, source in milestones:
    doc.add_heading(year, level=2)
    doc.add_paragraph("Significance: " + event)
    doc.add_paragraph("Source: " + source)

doc.add_paragraph(
    "Note: This is a historical evidence task, not a model-training exercise."
)

doc.save("AI_Historical_Timeline_Report.docx")

print("Word report created successfully!")